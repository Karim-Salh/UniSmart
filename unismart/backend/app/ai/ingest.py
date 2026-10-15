"""استيعاب المعرفة من لوحة الإدارة: نص / docx / xlsx / csv / sql ← مقاطع نصية أو سجلات جدولية.
أمان: ملفات SQL تُقرأ فقط (CREATE TABLE وINSERT)؛ أي أمر آخر يُتجاهل ولا يُنفَّذ أبدًا."""
import csv, io, re
def text_chunks(text: str, max_chars: int = 600) -> list[dict]:
    out, cur = [], ""
    for p in (x.strip() for x in re.split(r"\n+", text) if x.strip()):
        parts = re.split(r"(?<=[.!؟?])\s+", p) if len(p) > 700 else [p]
        for s in parts:
            if len(cur) + len(s) > max_chars and cur: out.append({"article": "", "content": cur.strip(), "meta": None}); cur = ""
            cur += ("\n" if cur else "") + s
    if cur.strip(): out.append({"article": "", "content": cur.strip(), "meta": None})
    return out[:3000]
def table_chunks(rows: list[list], table: str) -> list[dict]:
    if len(rows) < 2: return text_chunks("\n".join(" ".join(map(str, r)) for r in rows))
    hd = [str(h).strip() or f"عمود {i+1}" for i, h in enumerate(rows[0])]; out = []
    for n, r in enumerate(rows[1:5001], start=2):
        vals = [str(r[j]).strip() if j < len(r) and r[j] is not None else "" for j in range(len(hd))]
        txt = " | ".join(f"{h}: {v}" for h, v in zip(hd, vals) if v)
        if txt: out.append({"article": table, "content": txt, "meta": {"table": table, "cols": hd, "vals": vals, "row": n}})
    return out
def parse_docx(data: bytes) -> list[dict]:
    from docx import Document
    d = Document(io.BytesIO(data)); out = text_chunks("\n".join(p.text for p in d.paragraphs))
    for i, t in enumerate(d.tables, 1): out += table_chunks([[c.text for c in r.cells] for r in t.rows], f"جدول {i}")
    return out
def parse_xlsx(data: bytes) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [c for ws in wb.worksheets for c in table_chunks([list(r) for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)], ws.title)]
def parse_csv(text: str) -> list[dict]:
    try: dialect = csv.Sniffer().sniff(text[:2000], ",;\t")
    except csv.Error: dialect = csv.excel
    return table_chunks(list(csv.reader(io.StringIO(text), dialect)), "csv")
def _split(s: str, sep: str) -> list[str]:
    out, cur, depth, q = [], "", 0, None
    for i, ch in enumerate(s):
        if q:
            cur += ch
            if ch == q: q = None
        elif ch in "'\"`": q = ch; cur += ch
        elif ch == "(": depth += 1; cur += ch
        elif ch == ")": depth -= 1; cur += ch
        elif ch == sep and depth == 0: out.append(cur); cur = ""
        else: cur += ch
    if cur.strip(): out.append(cur)
    return out
_unq = lambda x: re.sub(r"^[`\"'\[]|[`\"'\]]$", "", x.strip()).replace("''", "'")
def parse_sql(text: str) -> tuple[list[dict], int]:
    schema, tables, skipped = {}, {}, 0
    text = re.sub(r"/\*.*?\*/", "", re.sub(r"--[^\n]*", "", text), flags=re.S)
    for st in (x.strip() for x in _split(text, ";") if x.strip()):
        m = re.match(r"create\s+table\s+(?:if\s+not\s+exists\s+)?([^\s(]+)\s*\((.*)\)[^)]*$", st, re.I | re.S)
        if m:
            schema[_unq(m[1]).lower()] = [_unq(p.split()[0]) for p in _split(m[2], ",") if not re.match(r"\s*(primary|foreign|unique|key|constraint|check|index)\b", p, re.I)]; continue
        m = re.match(r"insert\s+(?:ignore\s+)?into\s+([^\s(]+)\s*(?:\(([^)]*)\))?\s*values\s*(.*)$", st, re.I | re.S)
        if m:
            tn = _unq(m[1]).lower(); cols = [_unq(c) for c in m[2].split(",")] if m[2] else schema.get(tn, [])
            rows = tables.setdefault(tn, [cols])
            for tp in _split(m[3], ","): rows.append(["" if v.strip().lower() == "null" else _unq(v) for v in _split(tp.strip()[1:-1], ",")])
            continue
        skipped += 1
    return [c for tn, rows in tables.items() for c in table_chunks(rows, tn)], skipped
def parse(filename: str, data: bytes) -> tuple[list[dict], int]:
    n, skipped = filename.lower(), 0
    if n.endswith(".docx"): ch = parse_docx(data)
    elif n.endswith(".xlsx"): ch = parse_xlsx(data)
    elif n.endswith(".csv"): ch = parse_csv(data.decode("utf-8-sig", "ignore"))
    elif n.endswith(".sql"): ch, skipped = parse_sql(data.decode("utf-8", "ignore"))
    elif n.endswith((".txt", ".md")): ch = text_chunks(data.decode("utf-8", "ignore"))
    else: raise ValueError("نوع غير مدعوم؛ احفظ doc/xls القديمة بصيغة docx/xlsx")
    if not ch: raise ValueError("لم أجد محتوى قابلًا للفهرسة")
    return ch, skipped
