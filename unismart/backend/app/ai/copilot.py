"""المساعد المشترك للدكتور: توليد مسودات أسئلة من موضوع/شرائح (محليًا، استخراجيًا)، وتحليلات صفية. لا شيء يُنشر قبل اعتماد الدكتور."""
import re, random, collections
from sqlalchemy import select
from ..models import *
STOP = set("the of and to in is a for on with as by that this من في على الى إلى عن ما هو هي هذا هذه ذلك التي الذي ان أن او أو مع كل".split())
def extract_text(filename: str, data: bytes) -> str:
    n = filename.lower()
    if n.endswith(".pdf"):
        from pypdf import PdfReader; import io
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if n.endswith(".pptx"):
        from pptx import Presentation; import io
        return "\n".join(sh.text_frame.text for sl in Presentation(io.BytesIO(data)).slides for sh in sl.shapes if sh.has_text_frame)
    return data.decode("utf-8", "ignore")
def key_terms(text: str, n: int = 40) -> list[str]:
    freq = collections.Counter(w for w in re.findall(r"[\w\u0600-\u06FF]{4,}", text.lower()) if w not in STOP)
    return [w for w, _ in freq.most_common(n)]
def generate(text: str, n: int = 10, kinds=("mcq", "tf"), topic: str | None = None, seed: int = 7) -> list[dict]:
    rnd, terms, out = random.Random(seed), key_terms(text), []
    for s in (x.strip() for x in re.split(r"(?<=[.!؟?])\s+|\n+", text) if 40 <= len(x.strip()) <= 220):
        t = next((w for w in terms if re.search(rf"(?<!\w){re.escape(w)}(?!\w)", s.lower())), None)
        if not t or len(terms) < 5: continue
        stem = re.sub(re.escape(t), "_____", s, count=1, flags=re.I)
        if "mcq" in kinds:
            opts = rnd.sample([w for w in terms if w != t], 3) + [t]; rnd.shuffle(opts)
            out.append({"kind": "mcq", "stem": stem, "options": opts, "answer": opts.index(t), "topic": topic, "source_excerpt": s, "status": "DRAFT"})
        if "tf" in kinds and len(out) < n:
            fake = rnd.choice([w for w in terms if w != t]); true = rnd.random() < .5
            out.append({"kind": "tf", "stem": s if true else re.sub(re.escape(t), fake, s, count=1, flags=re.I), "options": ["صح", "خطأ"], "answer": 0 if true else 1, "topic": topic, "source_excerpt": s, "status": "DRAFT"})
        if len(out) >= n: break
    return out[:n]
def analytics(db, exam_id: int) -> dict:
    qs = db.scalars(select(Question).where(Question.exam_id == exam_id)).all()
    atts = db.scalars(select(Attempt).where(Attempt.exam_id == exam_id, Attempt.status.in_(["SUBMITTED", "AUTO_SUBMITTED"]))).all()
    ans = {(a.attempt_id, a.question_id): a.choice for a in db.scalars(select(Answer).where(Answer.attempt_id.in_([x.id for x in atts] or [0])))}
    ranked = sorted(atts, key=lambda a: a.score or 0); k = max(1, round(len(ranked) * .27)); low, high = ranked[:k], ranked[-k:]
    per_q, topics = [], collections.defaultdict(list)
    for q in qs:
        got = [ans.get((a.id, q.id)) for a in atts]; ok = [g == q.correct for g in got if g is not None]
        p = sum(ok) / len(ok) if ok else None
        wrong = collections.Counter(g for g in got if g is not None and g != q.correct)
        disc = (sum(ans.get((a.id, q.id)) == q.correct for a in high) - sum(ans.get((a.id, q.id)) == q.correct for a in low)) / k
        top = wrong.most_common(1)
        per_q.append({"question_id": q.id, "stem": q.text, "topic": q.topic, "p_correct": None if p is None else round(p, 2), "discrimination": round(disc, 2),
                      "common_mistake": {"option": q.options[top[0][0]], "share": round(top[0][1] / max(len(ok), 1), 2)} if top else None,
                      "flag": "راجع صياغة السؤال" if p is not None and (p < .3 or disc < .1) else None})
        if p is not None: topics[q.topic or "غير مصنّف"].append(p)
    mastery = sorted(({"topic": t, "mastery": round(sum(v) / len(v), 2)} for t, v in topics.items()), key=lambda x: x["mastery"])
    return {"students": len(atts), "avg_score": round(sum(a.score or 0 for a in atts) / max(len(atts), 1), 2), "weak_topics": mastery[:3], "mastery": mastery, "questions": per_q}
